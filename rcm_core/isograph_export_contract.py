"""Isograph RCM-Cost export layout contract (phase 0 SSOT for tests + import mapper)."""

from __future__ import annotations

from pathlib import Path
from typing import Mapping

import openpyxl

# All 25 sheets in AW export (structuur-fixture 2026-05).
EXPECTED_SHEET_NAMES: frozenset[str] = frozenset(
    {
        "RcmCauses",
        "Project",
        "RcmCauseEffectAssignments",
        "RcmScheduledTasks",
        "RcmCorrectiveTasks",
        "RcmLocations",
        "Labor",
        "Spares",
        "Equipments",
        "RcmScheduledTaskLabor",
        "RcmScheduledTaskEquipment",
        "RcmScheduledTaskSpares",
        "RcmCorrectiveTaskEquipment",
        "RcmCorrectiveTaskLabor",
        "RcmCorrectiveTaskSpares",
        "RcmProjectProfile",
        "RcmEffects",
        "RcmFunctionalFailures",
        "RcmFunctions",
        "TaskGroups",
        "TaskGroupTypes",
        "RcmEffectProfiles",
        "RcmEquipmentProfiles",
        "RcmLaborProfiles",
        "RcmSpareProfiles",
    }
)

# Minimale kolomheaders per must-v1 tabblad (import gate).
MUST_V1_SHEET_HEADERS: dict[str, tuple[str, ...]] = {
    "RcmLocations": ("Id", "Parent", "Description"),
    "RcmFunctions": ("Id", "Parent", "Description"),
    "RcmFunctionalFailures": ("Id", "Parent", "Description"),
    "RcmCauses": (
        "Id",
        "Parent",
        "Description",
        "FmMttf",
        "FmStd",
        "InitialAge",
        "Mttr",
        "FmDistribution",
        "LocationId",
    ),
    "RcmEffects": ("Id", "Description"),
    "RcmCauseEffectAssignments": (
        "Cause",
        "Effect",
        "RedundancyFactor",
        "PEnable",
        "IEnable",
        "CEnable",
    ),
    "RcmCorrectiveTasks": ("Cause", "TaskDuration", "Description"),
    "RcmScheduledTasks": (
        "Cause",
        "TaskDuration",
        "FixedInterval",
        "Enabled",
        "Type",
        "Description",
    ),
    "TaskGroups": ("Id", "Description"),
    "Project": ("LifeTime", "RcmNoSimulations", "RcmRandomNoSeed"),
}


def load_workbook_headers(path: Path) -> dict[str, tuple[str, ...]]:
    """Return first-row headers per sheet (trimmed, non-empty)."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        out: dict[str, tuple[str, ...]] = {}
        for name in wb.sheetnames:
            ws = wb[name]
            row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
            headers = tuple(
                str(c).strip()
                for c in (row or [])
                if c is not None and str(c).strip()
            )
            out[name] = headers
        return out
    finally:
        wb.close()


def assert_fixture_matches_contract(headers: Mapping[str, tuple[str, ...]]) -> None:
    """Raise AssertionError when workbook layout diverges from contract."""
    sheets = set(headers)
    if sheets != EXPECTED_SHEET_NAMES:
        missing = EXPECTED_SHEET_NAMES - sheets
        extra = sheets - EXPECTED_SHEET_NAMES
        raise AssertionError(
            f"Sheet mismatch: missing={sorted(missing)} extra={sorted(extra)}"
        )
    for sheet, required in MUST_V1_SHEET_HEADERS.items():
        present = set(headers[sheet])
        missing = [c for c in required if c not in present]
        if missing:
            raise AssertionError(f"{sheet} missing columns: {missing}")
