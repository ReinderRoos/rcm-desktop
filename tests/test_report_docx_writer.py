from __future__ import annotations

import json
from pathlib import Path

from docx import Document

from rcm_desktop.adapter.report_docx_writer import write_report_docx
from rcm_desktop.adapter.report_materialization_service import materialize_report
from rcm_desktop.adapter.report_options import ReportOptions
from rcm_desktop.adapter.report_run_source_service import ReportRunBundle, ReportScenarioRun
from rcm_desktop.adapter.planning_overlay_state import PlanningOverlayState
from rcm_desktop.adapter.run_service import run as run_single
from rcm_core.models import RCMProject


def _sample_project() -> RCMProject:
    raw = json.loads(Path("tests/fixtures/sample_project.rcm.json").read_text(encoding="utf-8"))
    return RCMProject.from_dict(raw)


def test_write_report_docx_creates_parseable_file(tmp_path: Path) -> None:
    project = _sample_project()
    run = run_single(project, Path("tests/fixtures/sample_project.rcm.json"))
    bundle = ReportRunBundle(
        mode="single",
        scenarios=(
            ReportScenarioRun(
                key="current",
                label="T",
                run_result=run,
                overlay_at_run=PlanningOverlayState.inactive(),
                scenario_key=None,
            ),
        ),
    )
    out = tmp_path / "rapport.docx"
    doc = materialize_report(
        project,
        bundle,
        ReportOptions(output_docx_path=out),
    )
    write_report_docx(doc, out)
    assert out.is_file()
    parsed = Document(str(out))
    assert len(parsed.paragraphs) + len(parsed.tables) >= 3
