"""Slice 70 issue 13 — regressie golden cause + CM-fixture in resultaten."""

from __future__ import annotations

import tempfile
from pathlib import Path

import pytest
from openpyxl import load_workbook

from rcm_core.effect_cost_parity import build_effect_cost_parity_report
from rcm_core.effect_impact_service import aggregate, effect_yearly_series
from rcm_core.incremental_run import run_incremental_analysis
from rcm_desktop.adapter.failure_validation_export_service import export_validation_excel
from rcm_desktop.adapter.fm_verification_service import build_fm_verification_view
from rcm_desktop.adapter.isograph_import_service import build_from_workbook
from rcm_desktop.adapter.run_service import build_run_result

CM_FIXTURE = Path(__file__).resolve().parent / "fixtures" / "RCMCostdata export_CM.xlsx"
GOLDEN_CAUSE = "06H-350.1.1.1.1.1.A.1"
GOLDEN_EFFECTS = (
    "VGM - Effect 3",
    "VGM - Effect 2",
    "Schutten 0-20% functieverlies",
    "Schutten 81-100% functieverlies",
)


@pytest.fixture(scope="module")
def cm_run():
    if not CM_FIXTURE.is_file():
        pytest.skip("CM fixture ontbreekt")
    result = build_from_workbook(CM_FIXTURE, modeljaar=2026)
    run_out = run_incremental_analysis(
        result.project,
        CM_FIXTURE.with_suffix(".rcm.cache.json"),
        full_recompute=True,
        parallel=False,
    )
    return result.project, build_run_result(result.project, list(run_out.fm_results.values()))


@pytest.mark.skipif(not CM_FIXTURE.is_file(), reason="CM fixture ontbreekt")
def test_golden_cause_four_rows_inspector_and_top10(cm_run) -> None:
    project, run = cm_run
    fmr = next(fr for fr in run.fm_core_results if fr.fm_id == GOLDEN_CAUSE)
    view = build_fm_verification_view(project, fmr)
    inspector_ids = {r.klasse_id for r in view.effect_result_rows}
    assert inspector_ids == set(GOLDEN_EFFECTS)

    rows = aggregate(
        project,
        {fmr.fm_id: fmr},
        fm_ids=frozenset({GOLDEN_CAUSE}),
    )
    row_ids = {r.klasse_id for r in rows}
    assert set(GOLDEN_EFFECTS) <= row_ids
    rf = {r.klasse_id: r.rf_display for r in rows if r.klasse_id in GOLDEN_EFFECTS}
    assert rf["VGM - Effect 3"] == "0,1"
    assert rf["VGM - Effect 2"] == "0,9"
    assert rf["Schutten 0-20% functieverlies"] == "1"
    assert rf["Schutten 81-100% functieverlies"] == "0,5"


@pytest.mark.skipif(not CM_FIXTURE.is_file(), reason="CM fixture ontbreekt")
def test_cm_fixture_pm_effect_visible_in_aggregation(cm_run) -> None:
    project, run = cm_run
    fmr = next(fr for fr in run.fm_core_results if fr.fm_id == GOLDEN_CAUSE)
    assert sum(fmr.pm_effect_bijdragen.values()) > 0.0
    view = build_fm_verification_view(project, fmr)
    assert any(r.waarde_pm > 0.0 for r in view.effect_result_rows)


@pytest.mark.skipif(not CM_FIXTURE.is_file(), reason="CM fixture ontbreekt")
def test_effect_yearly_series_reconciles_lifecycle(cm_run) -> None:
    project, run = cm_run
    fmr = next(fr for fr in run.fm_core_results if fr.fm_id == GOLDEN_CAUSE)
    klasse_id = "Schutten 0-20% functieverlies"
    series = effect_yearly_series(
        project,
        {fmr.fm_id: fmr},
        klasse_id,
        fm_ids=frozenset({GOLDEN_CAUSE}),
    )
    assert len(series) > 0
    lifecycle_total = next(
        r.waarde_totaal
        for r in build_fm_verification_view(project, fmr).effect_result_rows
        if r.klasse_id == klasse_id
    )
    assert sum(series) == pytest.approx(lifecycle_total, rel=0.05)


@pytest.mark.skipif(not CM_FIXTURE.is_file(), reason="CM fixture ontbreekt")
def test_validation_export_has_effect_tab(cm_run, tmp_path: Path) -> None:
    project, run = cm_run
    fm_map = {fr.fm_id: fr for fr in run.fm_core_results}
    out = tmp_path / "validation.xlsx"
    export_validation_excel(project, fm_map, out)
    wb = load_workbook(out, read_only=True)
    assert "Effectcategorieën" in wb.sheetnames


@pytest.mark.skipif(not CM_FIXTURE.is_file(), reason="CM fixture ontbreekt")
def test_effect_cost_parity_missing_benchmark_is_not_fail(cm_run) -> None:
    project, run = cm_run
    fm_map = {fr.fm_id: fr for fr in run.fm_core_results}
    report = build_effect_cost_parity_report(project, fm_map)
    assert report.rows
    assert all(r.verdict.value != "fail" for r in report.rows)
