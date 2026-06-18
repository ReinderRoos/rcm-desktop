"""Tests — failure validation counterfactuals + Excel export."""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook

from rcm_core.engine import run_analytical
from rcm_core.failure_parity_validation import build_failure_validation_report
from rcm_core.persistence import load_project
from rcm_desktop.adapter.failure_validation_export_service import (
    allocate_validation_causes,
    compute_validation_cell_fills,
    export_validation_excel,
    validation_column_headers,
)
from rcm_desktop.adapter import run_service
from rcm_desktop.adapter.isograph_import_service import build_from_workbook

CM_FIXTURE = Path(__file__).resolve().parent / "fixtures" / "RCMCostdata export_CM.xlsx"
GAARKEUKEN = Path(__file__).resolve().parent / "fixtures" / "RCMCostdata export_Gaarkeuken.rcm.json"


def test_validation_column_headers_include_counterfactuals() -> None:
    headers = validation_column_headers()
    assert "ef_cm_overlay" in headers
    assert "ef_horizon_forward" in headers
    assert "Dominante factor" in headers
    assert "Oorzaakcode(s)" in headers
    assert "Waarschijnlijke oorzaak" in headers
    assert "Actie" in headers


@pytest.mark.skipif(not CM_FIXTURE.is_file(), reason="CM fixture ontbreekt")
def test_build_failure_validation_report_on_cm_fixture() -> None:
    built = build_from_workbook(CM_FIXTURE, modeljaar=2026)
    fm_results, _ = run_analytical(built.project, parallel=False)
    report = build_failure_validation_report(built.project, fm_results)
    assert len(report.rows) == len(fm_results)
    assert report.cm_overlay_pm_count > 0
    row = report.rows[0]
    assert row.ef_actual >= 0.0
    assert row.ef_cm_overlay >= 0.0
    assert row.dominant_factor != ""


@pytest.mark.skipif(not GAARKEUKEN.is_file(), reason="Gaarkeuken fixture ontbreekt")
def test_export_validation_excel_gaarkeuken(tmp_path: Path) -> None:
    project = load_project(GAARKEUKEN)
    fm_results, _ = run_analytical(project, parallel=False)
    out = tmp_path / "validatie.xlsx"
    written = export_validation_excel(project, fm_results, out)
    assert written.is_file()

    wb = load_workbook(written)
    assert "Validatie" in wb.sheetnames
    assert "Legenda" in wb.sheetnames
    assert "Oorzaken" in wb.sheetnames
    ws = wb["Validatie"]
    headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    assert headers[0] == "FM"
    assert "ef_horizon_forward" in headers
    assert "Oorzaakcode(s)" in headers
    assert ws.auto_filter.ref is not None
    assert ws.freeze_panes == "D2"

    report = build_failure_validation_report(project, fm_results)
    failing = [
        r
        for r in report.rows
        if r.aw_total_w is not None and compute_validation_cell_fills(r).get("ef_actual")
    ]
    assert failing
    non_a1 = [
        r
        for r in failing
        if not (
            r.dominant_factor == "residu (AW-onbekend)"
            and "A1" in allocate_validation_causes(r).cause_codes
        )
    ]
    target = non_a1[0] if non_a1 else failing[0]
    alloc = allocate_validation_causes(target)
    assert alloc.likely_cause
    assert alloc.action
    ef_actual_col = headers.index("ef_actual (run)") + 1
    row_idx = next(i + 2 for i, r in enumerate(report.rows) if r.fm_id == target.fm_id)
    cell = ws.cell(row=row_idx, column=ef_actual_col)
    assert cell.fill.start_color.rgb in (
        "00FFC7CE",
        "FFFFC7CE",
        "00D9D9D9",
        "FFD9D9D9",
    )
    cause_col = headers.index("Oorzaakcode(s)") + 1
    cause_cell = ws.cell(row=row_idx, column=cause_col)
    assert cause_cell.value

    legend = wb["Legenda"]
    markering_found = any(
        row[0] == "Markeringen"
        for row in legend.iter_rows(min_row=1, values_only=True)
        if row and row[0]
    )
    assert markering_found
    wb.close()


@pytest.mark.skipif(not GAARKEUKEN.is_file(), reason="Gaarkeuken fixture ontbreekt")
def test_compute_validation_cell_fills_marks_parity_fail() -> None:
    project = load_project(GAARKEUKEN)
    fm_results, _ = run_analytical(project, parallel=False)
    report = build_failure_validation_report(project, fm_results)
    failing = [
        r
        for r in report.rows
        if r.aw_total_w is not None
        and abs(r.ef_actual - r.aw_total_w) > (r.aw_total_w_err or 0.01)
    ]
    assert failing, "verwacht minstens één FM met parity-fout t.o.v. AW TotalW"
    non_a1 = [
        r
        for r in failing
        if not (
            r.dominant_factor == "residu (AW-onbekend)"
            and "A1" in allocate_validation_causes(r).cause_codes
        )
    ]
    target = non_a1[0] if non_a1 else failing[0]
    fills = compute_validation_cell_fills(target)
    assert fills.get("ef_actual") is not None
    rgb = fills["ef_actual"].fgColor.rgb
    if non_a1:
        assert rgb in ("00FFC7CE", "FFFFC7CE")
    else:
        assert rgb in ("00FFC7CE", "FFFFC7CE", "00D9D9D9", "FFD9D9D9")


@pytest.mark.skipif(not GAARKEUKEN.is_file(), reason="Gaarkeuken fixture ontbreekt")
def test_a1_dominant_parity_fail_action_suggests_band_acceptance() -> None:
    """A1-dominante rijen krijgen band-accepteren actie, geen softwarefix-suggestie."""
    project = load_project(GAARKEUKEN)
    result = run_service.run(project, GAARKEUKEN, full_recompute=True)
    assert result.status == "done"
    fm_by_id = {fr.fm_id: fr for fr in result.fm_core_results}
    report = build_failure_validation_report(project, fm_by_id)

    a1_fails = [
        r
        for r in report.rows
        if r.aw_total_w is not None
        and "A1" in allocate_validation_causes(r).cause_codes
        and r.dominant_factor == "residu (AW-onbekend)"
    ]
    assert len(a1_fails) >= 10, "verwacht substantieel A1-residu op Gaarkeuken"
    for row in a1_fails[:5]:
        alloc = allocate_validation_causes(row)
        action_lower = alloc.action.lower()
        assert "accepteren" in action_lower or "band" in action_lower
        assert "softwarefix" in action_lower or "geen software" in action_lower


@pytest.mark.skipif(not GAARKEUKEN.is_file(), reason="Gaarkeuken fixture ontbreekt")
def test_b3_secondary_action_suggests_initial_age_harmonization() -> None:
    """B3-mismatch krijgt optionele leeftijd-harmonisatie hint in export-actie."""
    project = load_project(GAARKEUKEN)
    result = run_service.run(project, GAARKEUKEN, full_recompute=True)
    fm_by_id = {fr.fm_id: fr for fr in result.fm_core_results}
    report = build_failure_validation_report(project, fm_by_id)

    b3_rows = [
        r
        for r in report.rows
        if "B3" in allocate_validation_causes(r).cause_codes
    ]
    assert b3_rows, "verwacht minstens één B3-secundair signaal op Gaarkeuken"
    for row in b3_rows[:3]:
        action = allocate_validation_causes(row).action.lower()
        assert "harmoniseer leeftijd" in action or "initialage" in action


@pytest.mark.skipif(not GAARKEUKEN.is_file(), reason="Gaarkeuken fixture ontbreekt")
def test_a1_within_band_ef_actual_not_marked_red() -> None:
    """Parity OK (binnen AW-band): ef_actual krijgt groen, geen rood."""
    project = load_project(GAARKEUKEN)
    result = run_service.run(project, GAARKEUKEN, full_recompute=True)
    fm_by_id = {fr.fm_id: fr for fr in result.fm_core_results}
    report = build_failure_validation_report(project, fm_by_id)

    ok_rows = [
        r
        for r in report.rows
        if r.aw_total_w is not None
        and compute_validation_cell_fills(r).get("ef_actual") is not None
        and allocate_validation_causes(r).action == "Geen actie nodig"
    ]
    assert ok_rows, "verwacht minstens één Parity OK rij"
    fills = compute_validation_cell_fills(ok_rows[0])
    assert fills["ef_actual"].fgColor.rgb in ("00E2EFDA", "FFE2EFDA")

