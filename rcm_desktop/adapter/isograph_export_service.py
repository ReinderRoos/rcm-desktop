"""RCM-Cost / AW workbook export via bron-sidecar patch (slice 77, ADR-0011)."""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.worksheet.worksheet import Worksheet

from rcm_core.import_settings_contract import SOURCE_WORKBOOK_PATH_KEY
from rcm_core.models import FailureType, RCMProject

from rcm_desktop.adapter.isograph_import_service import HOURS_PER_YEAR
from rcm_desktop.adapter.project_path_resolution_service import ResolvedProjectPath


@dataclass
class ExportSheetCounts:
    patched: int = 0
    added: int = 0
    warned: int = 0


@dataclass
class IsographExportResult:
    output_path: Path
    counts_by_sheet: dict[str, ExportSheetCounts] = field(default_factory=dict)
    warnings: tuple[str, ...] = ()

    @property
    def total_patched(self) -> int:
        return sum(c.patched for c in self.counts_by_sheet.values())

    @property
    def total_added(self) -> int:
        return sum(c.added for c in self.counts_by_sheet.values())

    @property
    def total_warned(self) -> int:
        return sum(c.warned for c in self.counts_by_sheet.values())


def resolve_source_workbook_path(
    project_file_path: Path,
    import_settings: dict[str, Any],
) -> Path | None:
    candidates: list[Path] = []
    key = import_settings.get(SOURCE_WORKBOOK_PATH_KEY)
    if key:
        raw = Path(str(key))
        candidates.append(raw)
        if not raw.is_absolute():
            candidates.append(project_file_path.parent / raw.name)
    default = ResolvedProjectPath(project_file_path).source_workbook_path()
    if default is not None:
        candidates.append(default)
    for candidate in candidates:
        if candidate.is_file():
            return candidate
    return None


def export_project_to_workbook(
    project: RCMProject,
    source_workbook_path: Path,
    output_path: Path,
) -> IsographExportResult:
    wb = openpyxl.load_workbook(source_workbook_path)
    counts: dict[str, ExportSheetCounts] = {}
    warnings: list[str] = []

    if "RcmCauses" in wb.sheetnames:
        sheet_counts, sheet_warnings = _patch_rcm_causes(
            wb["RcmCauses"],
            project,
        )
        counts["RcmCauses"] = sheet_counts
        warnings.extend(sheet_warnings)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return IsographExportResult(
        output_path=output_path,
        counts_by_sheet=counts,
        warnings=tuple(warnings),
    )


def _years_to_hours(years: float) -> float:
    return float(years) * HOURS_PER_YEAR


def _failure_distribution(failure_type: FailureType) -> str:
    if failure_type == FailureType.AGING:
        return "Normal"
    return "Exponential"


def _header_columns(ws: Worksheet) -> dict[str, int]:
    columns: dict[str, int] = {}
    for col in range(1, ws.max_column + 1):
        raw = ws.cell(1, col).value
        if raw is None:
            continue
        name = str(raw).strip()
        if name:
            columns[name] = col
    return columns


def _set_cell(ws: Worksheet, row: int, columns: dict[str, int], header: str, value: Any) -> bool:
    col = columns.get(header)
    if col is None:
        return False
    current = ws.cell(row, col).value
    if current == value:
        return False
    ws.cell(row, col, value=value)
    return True


def _patch_rcm_causes(
    ws: Worksheet,
    project: RCMProject,
) -> tuple[ExportSheetCounts, list[str]]:
    counts = ExportSheetCounts()
    warnings: list[str] = []
    columns = _header_columns(ws)
    if "Id" not in columns:
        return counts, warnings

    id_col = columns["Id"]
    seen_project_ids = set(project.faalwijzes)
    matched_rows: set[str] = set()

    for row in range(2, ws.max_row + 1):
        fm_id = str(ws.cell(row, id_col).value or "").strip()
        if not fm_id:
            continue
        fm = project.faalwijzes.get(fm_id)
        if fm is None:
            counts.warned += 1
            warnings.append(
                f"RcmCauses: rij '{fm_id}' ontbreekt in RCM2-project (niet verwijderd)"
            )
            continue
        matched_rows.add(fm_id)
        patched = False
        pbs = project.pbs_items.get(fm.pbs_id)
        initial_age_h = 0.0
        if pbs is not None:
            initial_age_h = _years_to_hours(pbs.current_age(int(project.config.modeljaar)))
        patch_values: dict[str, Any] = {
            "Description": fm.faalwijze_omschrijving,
            "LocationId": fm.pbs_id,
            "FmMttf": _years_to_hours(fm.mttf_jaar),
            "FmStd": _years_to_hours(fm.sigma_jaar),
            "InitialAge": initial_age_h,
            "Mttr": float(fm.downtime_per_failure.to_hours()),
            "FmDistribution": _failure_distribution(fm.failure_type),
        }
        for header, value in patch_values.items():
            if _set_cell(ws, row, columns, header, value):
                patched = True
        if patched:
            counts.patched += 1

    for fm_id in sorted(seen_project_ids - matched_rows):
        counts.added += 1
        warnings.append(
            f"RcmCauses: '{fm_id}' ontbreekt in bron — append nog niet geïmplementeerd"
        )

    return counts, warnings
