"""Read Isograph RCM-Cost export workbooks into sheet row dicts (Qt-free)."""

from __future__ import annotations

from pathlib import Path
from typing import Any

import openpyxl

from rcm_core.isograph_export_contract import MUST_V1_SHEET_HEADERS


def read_workbook_sheets(path: Path) -> dict[str, list[dict[str, Any]]]:
    """Load must-v1 sheets as list of row dicts (header → value)."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        out: dict[str, list[dict[str, Any]]] = {}
        for sheet_name in MUST_V1_SHEET_HEADERS:
            if sheet_name not in wb.sheetnames:
                continue
            ws = wb[sheet_name]
            rows = ws.iter_rows(values_only=True)
            header_row = next(rows, None)
            if not header_row:
                out[sheet_name] = []
                continue
            headers = [str(c).strip() if c is not None else "" for c in header_row]
            data: list[dict[str, Any]] = []
            for row in rows:
                if row is None or not any(c is not None and str(c).strip() for c in row):
                    continue
                values = list(row)
                if len(values) < len(headers):
                    values.extend([None] * (len(headers) - len(values)))
                record = {
                    headers[i]: values[i]
                    for i in range(len(headers))
                    if headers[i]
                }
                if record.get("Id") or record.get("Cause") or record.get("LifeTime") is not None:
                    data.append(record)
                elif sheet_name == "Project" and record:
                    data.append(record)
            out[sheet_name] = data
        return out
    finally:
        wb.close()
