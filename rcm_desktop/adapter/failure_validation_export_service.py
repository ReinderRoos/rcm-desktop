"""Validatie-export — Excel per faalwijze met counterfactual # falen (slice 65+)."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Mapping, Sequence

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from rcm_core.effect_impact_service import aggregate
from rcm_core.failure_parity_validation import (
    FailureValidationReport,
    FailureValidationRow,
    build_failure_validation_report,
)
from rcm_core.models import FMResult, RCMProject


@dataclass(frozen=True)
class ValidationExportColumn:
    header: str
    key: str


@dataclass(frozen=True)
class CauseAllocation:
    """Toegewezen oorzaak(en) voor een FM-validatierij (ADR-0008 v1.2)."""

    cause_codes: tuple[str, ...]
    likely_cause: str
    action: str


_VALIDATION_COLUMNS: tuple[ValidationExportColumn, ...] = (
    ValidationExportColumn("FM", "fm_id"),
    ValidationExportColumn("Omschrijving", "omschrijving"),
    ValidationExportColumn("failure_type", "failure_type"),
    ValidationExportColumn("AW TotalW", "aw_total_w"),
    ValidationExportColumn("AW TotalWErr", "aw_total_w_err"),
    ValidationExportColumn("AW OutageFrequency", "aw_outage_frequency"),
    ValidationExportColumn("AW InitialAge jr", "aw_initial_age_years"),
    ValidationExportColumn("AW FmMttf jr", "aw_mttf_years"),
    ValidationExportColumn("RCM2 leeftijd jr", "current_age_years"),
    ValidationExportColumn("RCM2 MTTF jr", "mttf_years"),
    ValidationExportColumn("RCM2 repair_quality", "repair_quality"),
    ValidationExportColumn("RCM2 multipliciteit", "effective_multiplicity"),
    ValidationExportColumn("REV actief", "rev_moments_active"),
    ValidationExportColumn("REV structureel", "rev_moments_structural"),
    ValidationExportColumn("CM-overlay # PM uit", "cm_overlay_disabled_pm_count"),
    ValidationExportColumn("ef_actual (run)", "ef_actual"),
    ValidationExportColumn("ef_cm_overlay", "ef_cm_overlay"),
    ValidationExportColumn("ef_no_rev", "ef_no_rev"),
    ValidationExportColumn("ef_rev_aw_100pct", "ef_rev_aw_100pct"),
    ValidationExportColumn("ef_rq_0", "ef_rq_0"),
    ValidationExportColumn("ef_rq_0.5", "ef_rq_0_5"),
    ValidationExportColumn("ef_rq_1", "ef_rq_1"),
    ValidationExportColumn("ef_horizon_forward", "ef_horizon_forward"),
    ValidationExportColumn("Δ scenario", "delta_scenario"),
    ValidationExportColumn("Δ REV effect", "delta_rev_effect"),
    ValidationExportColumn("Δ REV aging 100%", "delta_rev_pct"),
    ValidationExportColumn("Δ horizon", "delta_horizon"),
    ValidationExportColumn("Δ residu (AW−horizon)", "delta_residual"),
    ValidationExportColumn("Dominante factor", "dominant_factor"),
    ValidationExportColumn("Oorzaakcode(s)", "cause_codes"),
    ValidationExportColumn("Waarschijnlijke oorzaak", "likely_cause"),
    ValidationExportColumn("Actie", "recommended_action"),
)

_LEGEND_ROWS: tuple[tuple[str, str], ...] = (
    ("Doel", "Per FM isoleren waarom RCM2 # falen afwijkt van AW TotalW."),
    (
        "ef_actual",
        "Verwachte falen uit de huidige run (parity-dialoog).",
    ),
    (
        "ef_cm_overlay",
        "Analytisch herberekend met AW CM-overlay (aw_disabled_pm_ids).",
    ),
    (
        "ef_no_rev",
        "Zelfde als cm_overlay maar zonder REV-schema (alleen aging).",
    ),
    (
        "ef_rev_aw_100pct",
        "CM-overlay + REV met aging_effect_pct=100% (AW-default voor REV).",
    ),
    (
        "ef_rq_*",
        "Repair-quality sweep (0=as-old, 1=as-new) onder CM-overlay.",
    ),
    (
        "ef_horizon_forward",
        "Studieduur = volledige LifeTime vooruit (AW MC-semantiek).",
    ),
    ("Δ scenario", "ef_actual − ef_cm_overlay → PM/REV-scenario mismatch."),
    ("Δ REV effect", "ef_cm_overlay − ef_no_rev → REV in motor."),
    ("Δ REV aging 100%", "ef_rev_aw_100pct − ef_no_rev → ontbrekende aging_effect import."),
    ("Δ horizon", "ef_horizon_forward − ef_cm_overlay → LifeTime-semantiek."),
    ("Δ residu", "AW TotalW − ef_horizon_forward → rest (MC, Quantity, …)."),
    (
        "Dominante factor",
        "Grootste |Δ| onder scenario, REV, aging, horizon, residu.",
    ),
    (
        "Oorzaakcode(s)",
        "Codes A–F uit oorzakenmatrix (primair + secundaire invoer/scenario-signalen).",
    ),
    (
        "Waarschijnlijke oorzaak",
        "Korte verklaring op basis van counterfactuals, Δ's en invoerhints.",
    ),
    (
        "Actie",
        "Aanbevolen vervolgstap om parity te verbeteren of te accepteren.",
    ),
)

_CAUSE_CATALOG_ROWS: tuple[tuple[str, str, str], ...] = (
    ("A1", "Analytisch vs Monte Carlo", "Accepteren binnen AW-band (TotalW ± TotalWErr); geen softwarefix"),
    ("A2", "LifeTime-semantiek", "Productkeuze: AW-horizon vs resterende studieduur"),
    ("A3", "Verdelingsimplementatie", "Verdeling/parameters handmatig vergelijken (FmStd/beta)"),
    ("A4", "Random falen", "Focus op MTTF, multipliciteit, MC-residu"),
    ("B1", "Repair quality niet geïmporteerd", "AW repair quality importeren / handmatig zetten"),
    ("B2", "REV aging_effect_pct niet geïmporteerd", "aging_effect_pct=100% op REV-taken uit AW"),
    ("B3", "InitialAge vs RCM2-leeftijd", "Bouwjaar/InitialAge-mapping controleren"),
    ("B4", "MTTF/sigma/verdeling invoer", "FmMttf, FmStd, FmDistribution controleren"),
    ("B5", "Multipliciteit / Quantity", "Quantity/PBS-keten vs AW vergelijken"),
    ("B6", "TotalW vs OutageFrequency benchmark", "TotalW als AW-bron gebruiken"),
    ("C1", "Run ≠ AW CM-scenario", "Run met CM-overlay materialiseren"),
    ("C2", "REV uit in AW CM-scenario", "Overlay respecteren; niet extra REV activeren"),
    ("C3", "PM-import fout", "ScheduledTasks/import controleren"),
    ("C4", "Handmatige planning-overlay", "Overlay resetten naar import-seed"),
    ("D1", "REV-interval / momenten", "Intervallen (TaskInterval) controleren"),
    ("D2", "REV-motor vs AW", "REV-algoritme vergelijken met AW"),
    ("D4", "REV op random FM (n.v.t.)", "Geen REV-actie; focus op MTTF/horizon/residu"),
    ("E1", "Repair quality waarde", "Juiste repair_quality instellen (zie ef_rq_*)"),
    ("F1", "MC-sampling variance", "Accepteren binnen TotalWErr-band"),
    ("F*", "Overig residu (sigma, NMF, …)", "Handmatig dieper onderzoeken"),
)

_MARKING_LEGEND_ROWS: tuple[tuple[str, str, str], ...] = (
    ("Rood", "FFC7CE", "Parity-fout: ef_actual buiten AW-band (TotalW ± TotalWErr)."),
    (
        "Oranje",
        "FFEB9C",
        "Invoer-/importsuggestie: leeftijd/MTTF verschilt, repair_quality of REV-scenario.",
    ),
    (
        "Geel",
        "FFF2CC",
        "Significante Δ (attribuut); donkerder geel = dominante factor bij parity-fout.",
    ),
    (
        "Groen",
        "C6EFCE",
        "Counterfactual het dichtst bij AW TotalW (handvat voor verbetering).",
    ),
    (
        "Grijs",
        "D9D9D9",
        "A1 model-keuze: analytisch vs MC — vergelijk met AW-band; geen softwarefix.",
    ),
    (
        "Blauw",
        "DDEBF7",
        "Dominante factor-kolom wanneer er een AW-benchmark is en |Δ residu| > band.",
    ),
)
_FILL_PARITY_FAIL = PatternFill(fill_type="solid", fgColor="FFC7CE")
_FILL_PARITY_OK = PatternFill(fill_type="solid", fgColor="E2EFDA")
_FILL_A1_ACCEPT = PatternFill(fill_type="solid", fgColor="D9D9D9")
_FILL_INPUT_WARN = PatternFill(fill_type="solid", fgColor="FFEB9C")
_FILL_DELTA = PatternFill(fill_type="solid", fgColor="FFF2CC")
_FILL_DELTA_DOMINANT = PatternFill(fill_type="solid", fgColor="FFE699")
_FILL_CLOSEST = PatternFill(fill_type="solid", fgColor="C6EFCE")
_FILL_DOMINANT_COL = PatternFill(fill_type="solid", fgColor="DDEBF7")

_INPUT_YEARS_TOL = 0.05
_INPUT_MTTF_REL_TOL = 0.02

_DELTA_KEYS: frozenset[str] = frozenset(
    {
        "delta_scenario",
        "delta_rev_effect",
        "delta_rev_pct",
        "delta_horizon",
        "delta_residual",
    }
)

_COUNTERFACTUAL_KEYS: tuple[str, ...] = (
    "ef_actual",
    "ef_cm_overlay",
    "ef_no_rev",
    "ef_rev_aw_100pct",
    "ef_rq_0",
    "ef_rq_0_5",
    "ef_rq_1",
    "ef_horizon_forward",
)

_DOMINANT_FACTOR_TO_DELTA_KEY: dict[str, str] = {
    "scenario (CM-overlay)": "delta_scenario",
    "REV in motor": "delta_rev_effect",
    "REV aging_effect import": "delta_rev_pct",
    "LifeTime-semantiek": "delta_horizon",
    "residu (AW-onbekend)": "delta_residual",
}


def _failure_tolerance(row: FailureValidationRow) -> float:
    aw = row.aw_total_w
    if aw is None:
        return 0.0
    err = row.aw_total_w_err
    if err is not None and float(err) > 0:
        return float(err)
    return max(0.01, abs(float(aw)) * 0.05)


def _parity_gap(row: FailureValidationRow) -> float | None:
    if row.aw_total_w is None:
        return None
    return row.ef_actual - float(row.aw_total_w)


def _is_parity_fail(row: FailureValidationRow) -> bool:
    gap = _parity_gap(row)
    if gap is None:
        return False
    tol = _failure_tolerance(row)
    if tol <= 0:
        return max(abs(row.aw_total_w or 0.0), row.ef_actual, 0.01) * 0.05 < abs(gap)
    return abs(gap) > tol


def _is_parity_ok(row: FailureValidationRow) -> bool:
    return row.aw_total_w is not None and not _is_parity_fail(row)


def _years_mismatch(a: float | None, b: float | None) -> bool:
    if a is None or b is None:
        return False
    return abs(float(a) - float(b)) > _INPUT_YEARS_TOL


def _mttf_mismatch(row: FailureValidationRow) -> bool:
    aw = row.aw_mttf_years
    rcm = row.mttf_years
    if aw is None or aw <= 0:
        return False
    return abs(rcm - aw) / aw > _INPUT_MTTF_REL_TOL


def _rev_scenario_mismatch(row: FailureValidationRow) -> bool:
    return row.rev_moments_structural > 0 and row.rev_moments_active < row.rev_moments_structural


def _repair_quality_import_hint(row: FailureValidationRow) -> bool:
    if row.aw_total_w is None or row.repair_quality < 0.99:
        return False
    tol = _failure_tolerance(row)
    cm_dist = abs(row.ef_cm_overlay - row.aw_total_w)
    rq0_dist = abs(row.ef_rq_0 - row.aw_total_w)
    return rq0_dist + tol < cm_dist


def _rev_aging_import_hint(row: FailureValidationRow) -> bool:
    if row.failure_type != "aging" or row.aw_total_w is None:
        return False
    gain = row.delta_rev_pct
    if gain is None:
        return False
    tol = _failure_tolerance(row)
    if abs(gain) <= tol:
        return False
    return abs(row.ef_cm_overlay - row.ef_no_rev) <= tol


def _closest_counterfactual_key(row: FailureValidationRow) -> str | None:
    if row.aw_total_w is None:
        return None
    best_key: str | None = None
    best_dist = float("inf")
    for key in _COUNTERFACTUAL_KEYS:
        val = float(getattr(row, key))
        dist = abs(val - row.aw_total_w)
        if dist < best_dist:
            best_dist = dist
            best_key = key
    return best_key


def _dominant_delta_key(row: FailureValidationRow) -> str | None:
    return _DOMINANT_FACTOR_TO_DELTA_KEY.get(row.dominant_factor)


def _is_closest_counterfactual(row: FailureValidationRow, key: str) -> bool:
    closest = _closest_counterfactual_key(row)
    return closest == key and row.aw_total_w is not None


def _residu_allocation(row: FailureValidationRow, tol: float) -> CauseAllocation:
    residual = row.delta_residual
    if residual is not None and abs(float(residual)) <= tol:
        return CauseAllocation(
            ("F1",),
            "MC-sampling variance (binnen TotalWErr)",
            "Accepteren binnen band",
        )
    if row.failure_type == "random":
        return CauseAllocation(
            ("A4", "A1"),
            "Random falen + analytisch/MC-residu",
            "MTTF/multipliciteit controleren; rest via MC-band",
        )
    return CauseAllocation(
        ("A1", "F*"),
        "Analytisch vs Monte Carlo / onbekend residu",
        "Accepteren: vergelijk met AW-band (TotalW ± TotalWErr); geen softwarefix",
    )


def _is_a1_dominant_model_choice(row: FailureValidationRow, allocation: CauseAllocation) -> bool:
    return (
        row.dominant_factor == "residu (AW-onbekend)"
        and bool(allocation.cause_codes)
        and allocation.cause_codes[0] == "A1"
    )


def _optional_b3_action(secondary: tuple[str, ...]) -> str | None:
    if "B3" not in secondary:
        return None
    return "Optioneel: harmoniseer leeftijd met AW InitialAge (invoer)"


def _merge_secondary_action(primary_action: str, secondary: tuple[str, ...]) -> str:
    b3 = _optional_b3_action(secondary)
    if b3 is None:
        return primary_action
    if primary_action.startswith("Optioneel"):
        return primary_action
    return f"{primary_action}; {b3}"


def allocate_validation_causes(row: FailureValidationRow) -> CauseAllocation:
    """Wijs oorzaakcode(s) toe op basis van counterfactuals, Δ's en invoerhints."""
    if row.aw_total_w is None:
        if row.aw_outage_frequency is not None:
            return CauseAllocation(
                ("B6",),
                "Benchmark-definitie AW (TotalW ontbreekt)",
                "TotalW als bron gebruiken i.p.v. OutageFrequency-afleiding",
            )
        return CauseAllocation(
            ("—",),
            "Geen AW-benchmark beschikbaar",
            "AW TotalW in import beschikbaar maken",
        )

    tol = _failure_tolerance(row)
    secondary: list[str] = []

    if _years_mismatch(row.aw_initial_age_years, row.current_age_years):
        secondary.append("B3")
    if _mttf_mismatch(row):
        secondary.append("B4")
    if _rev_scenario_mismatch(row):
        secondary.append("C2")

    if not _is_parity_fail(row):
        codes = tuple(dict.fromkeys(secondary))
        if codes:
            b3_only = codes == ("B3",)
            return CauseAllocation(
                codes,
                "Parity OK; secundaire invoer/scenario-afwijkingen",
                _optional_b3_action(codes)
                if b3_only
                else "Optioneel invoer harmoniseren (leeftijd/MTTF/REV-scenario)",
            )
        return CauseAllocation((), "Parity OK", "Geen actie nodig")

    primary_code: str
    likely: str
    action: str

    if _is_closest_counterfactual(row, "ef_horizon_forward"):
        primary_code = "A2"
        likely = "LifeTime-semantiek (resterende studieduur vs volledige LifeTime)"
        action = "Productkeuze: AW-horizon vs resterende studieduur"
    elif _is_closest_counterfactual(row, "ef_rq_0") and _repair_quality_import_hint(row):
        primary_code = "B1"
        likely = "Repair quality niet geïmporteerd (AW ≈ as-old)"
        action = "AW repair quality importeren / handmatig zetten"
    elif _is_closest_counterfactual(row, "ef_rq_0_5") and _repair_quality_import_hint(row):
        primary_code = "E1"
        likely = "Repair quality waarde wijkt af (tussen as-old en as-new)"
        action = "Juiste repair_quality instellen (zie ef_rq_*-sweep)"
    elif _is_closest_counterfactual(row, "ef_rev_aw_100pct") and _rev_aging_import_hint(row):
        primary_code = "B2"
        likely = "REV aging_effect_pct niet geïmporteerd (AW ≈ 100%)"
        action = "aging_effect_pct=100% op REV-taken uit AW"
    elif _is_closest_counterfactual(row, "ef_cm_overlay"):
        primary_code = "C1"
        likely = "Run wijkt af van AW CM-overlay-scenario"
        action = "Run met CM-overlay materialiseren"
    elif row.dominant_factor == "scenario (CM-overlay)":
        primary_code = "C4" if row.cm_overlay_disabled_pm_count > 0 else "C1"
        likely = (
            "Handmatige planning-overlay t.o.v. import-seed"
            if primary_code == "C4"
            else "Run ≠ AW CM-scenario"
        )
        action = (
            "Overlay resetten naar import-seed"
            if primary_code == "C4"
            else "Run met CM-overlay materialiseren"
        )
    elif row.dominant_factor == "REV aging_effect import":
        primary_code = "B2"
        likely = "REV aging_effect_pct niet geïmporteerd"
        action = "aging_effect_pct=100% op REV-taken uit AW"
    elif row.dominant_factor == "REV in motor":
        if _rev_scenario_mismatch(row):
            primary_code = "C2"
            likely = "REV uitgeschakeld in AW CM, actief in RCM2"
            action = "Overlay respecteren; niet extra REV activeren"
        elif row.failure_type != "aging":
            primary_code = "D4"
            likely = "REV-kolommen irrelevant bij random falen"
            action = "Geen REV-actie; focus op MTTF/horizon/residu"
        else:
            primary_code = "D1"
            likely = "REV-planning / REV-effect in motor"
            action = "Intervallen (TaskInterval) en REV-taken controleren"
    elif row.dominant_factor == "LifeTime-semantiek":
        primary_code = "A2"
        likely = "LifeTime-semantiek (horizon)"
        action = "Productkeuze: AW-horizon vs resterende studieduur"
    elif row.dominant_factor == "residu (AW-onbekend)":
        residu = _residu_allocation(row, tol)
        primary_code = residu.cause_codes[0]
        likely = residu.likely_cause
        action = residu.action
    else:
        residu = _residu_allocation(row, tol)
        primary_code = residu.cause_codes[0]
        likely = residu.likely_cause
        action = residu.action

    codes = tuple(dict.fromkeys([primary_code, *secondary]))
    return CauseAllocation(
        cause_codes=codes,
        likely_cause=likely,
        action=_merge_secondary_action(action, tuple(c for c in codes if c != primary_code)),
    )


def _cause_codes_text(allocation: CauseAllocation) -> str:
    return ", ".join(allocation.cause_codes)


def _apply_fill(
    fills: dict[str, PatternFill],
    key: str,
    fill: PatternFill,
    *,
    protected: frozenset[str] = frozenset(),
) -> None:
    if key in protected:
        return
    fills[key] = fill


def compute_validation_cell_fills(row: FailureValidationRow) -> dict[str, PatternFill]:
    """Per kolom-key een achtergrondkleur voor afwijkingen (pytestbaar)."""
    fills: dict[str, PatternFill] = {}
    tol = _failure_tolerance(row)
    parity_fail = _is_parity_fail(row)
    parity_ok = _is_parity_ok(row)
    closest = _closest_counterfactual_key(row)
    dominant_delta = _dominant_delta_key(row)
    locked: set[str] = set()

    if parity_fail:
        allocation = allocate_validation_causes(row)
        if _is_a1_dominant_model_choice(row, allocation):
            fills["ef_actual"] = _FILL_A1_ACCEPT
            locked.add("ef_actual")
        else:
            fills["ef_actual"] = _FILL_PARITY_FAIL
            fills["aw_total_w"] = _FILL_PARITY_FAIL
            locked.update(("ef_actual", "aw_total_w"))
    elif parity_ok:
        fills["ef_actual"] = _FILL_PARITY_OK
        locked.add("ef_actual")

    if _years_mismatch(row.aw_initial_age_years, row.current_age_years):
        for key in ("aw_initial_age_years", "current_age_years"):
            _apply_fill(fills, key, _FILL_INPUT_WARN, protected=frozenset(locked))
            locked.add(key)

    if _mttf_mismatch(row):
        for key in ("aw_mttf_years", "mttf_years"):
            _apply_fill(fills, key, _FILL_INPUT_WARN, protected=frozenset(locked))
            locked.add(key)

    if _rev_scenario_mismatch(row):
        for key in ("rev_moments_active", "rev_moments_structural"):
            _apply_fill(fills, key, _FILL_INPUT_WARN, protected=frozenset(locked))
            locked.add(key)

    if _repair_quality_import_hint(row):
        for key in ("repair_quality", "ef_rq_0"):
            _apply_fill(fills, key, _FILL_INPUT_WARN, protected=frozenset(locked))
            locked.add(key)

    if _rev_aging_import_hint(row):
        _apply_fill(fills, "ef_rev_aw_100pct", _FILL_INPUT_WARN, protected=frozenset(locked))
        locked.add("ef_rev_aw_100pct")

    for key in _DELTA_KEYS:
        raw = getattr(row, key)
        if raw is None or abs(float(raw)) <= tol:
            continue
        delta_fill = (
            _FILL_DELTA_DOMINANT if parity_fail and key == dominant_delta else _FILL_DELTA
        )
        _apply_fill(fills, key, delta_fill, protected=frozenset(locked))
        locked.add(key)

    if (
        row.aw_total_w is not None
        and row.delta_residual is not None
        and abs(float(row.delta_residual)) > tol
    ):
        _apply_fill(fills, "dominant_factor", _FILL_DOMINANT_COL, protected=frozenset(locked))

    if closest is not None:
        if parity_fail and closest != "ef_actual":
            _apply_fill(fills, closest, _FILL_CLOSEST, protected=frozenset(locked))
        elif parity_ok:
            _apply_fill(fills, closest, _FILL_CLOSEST, protected=frozenset(locked))

    if parity_fail:
        allocation = allocate_validation_causes(row)
        if allocation.cause_codes:
            for key in ("cause_codes", "likely_cause", "recommended_action"):
                _apply_fill(fills, key, _FILL_DOMINANT_COL, protected=frozenset(locked))

    return fills


def build_validation_report(
    project: RCMProject,
    fm_results: Mapping[str, FMResult],
) -> FailureValidationReport:
    return build_failure_validation_report(project, fm_results)


def _cell_value(row: FailureValidationRow, key: str, *, allocation: CauseAllocation | None = None) -> object:
    alloc = allocation or allocate_validation_causes(row)
    if key == "cause_codes":
        return _cause_codes_text(alloc)
    if key == "likely_cause":
        return alloc.likely_cause
    if key == "recommended_action":
        return alloc.action
    value = getattr(row, key)
    if isinstance(value, float):
        return value
    if value is None:
        return ""
    return value


def _autosize_columns(ws, *, max_width: int = 48) -> None:
    for col_idx, column in enumerate(_VALIDATION_COLUMNS, start=1):
        letter = get_column_letter(col_idx)
        header_len = len(column.header)
        max_len = header_len
        for row_idx in range(2, ws.max_row + 1):
            cell = ws[f"{letter}{row_idx}"]
            if cell.value is not None:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[letter].width = min(max_width, max_len + 2)


def export_validation_workbook(
    report: FailureValidationReport,
    path: Path | str,
    *,
    projectnaam: str = "",
    input_file: Path | str | None = None,
    project: RCMProject | None = None,
    fm_results: Mapping[str, FMResult] | None = None,
) -> Path:
    """Schrijf validatie-Excel naar path; retourneert opgelost pad."""
    dest = Path(path)
    wb = Workbook()

    ws = wb.active
    ws.title = "Validatie"
    bold = Font(bold=True)
    for col_idx, column in enumerate(_VALIDATION_COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=column.header)
        cell.font = bold

    data_end_row = len(report.rows) + 1
    for row_idx, row in enumerate(report.rows, start=2):
        allocation = allocate_validation_causes(row)
        row_fills = compute_validation_cell_fills(row)
        for col_idx, column in enumerate(_VALIDATION_COLUMNS, start=1):
            cell = ws.cell(
                row=row_idx,
                column=col_idx,
                value=_cell_value(row, column.key, allocation=allocation),
            )
            fill = row_fills.get(column.key)
            if fill is not None:
                cell.fill = fill

    if report.rows:
        last_col = get_column_letter(len(_VALIDATION_COLUMNS))
        ws.auto_filter.ref = f"A1:{last_col}{data_end_row}"
        ws.freeze_panes = "D2"

    meta_row = len(report.rows) + 3
    ws.cell(row=meta_row, column=1, value="Geëxporteerd").font = bold
    ws.cell(row=meta_row, column=2, value=datetime.now(timezone.utc).isoformat())
    ws.cell(row=meta_row + 1, column=1, value="Inputbestand").font = bold
    ws.cell(row=meta_row + 1, column=2, value=str(Path(input_file).name) if input_file else "")
    ws.cell(row=meta_row + 2, column=1, value="Project").font = bold
    ws.cell(row=meta_row + 2, column=2, value=projectnaam)
    ws.cell(row=meta_row + 3, column=1, value="LifeTime jr").font = bold
    ws.cell(row=meta_row + 3, column=2, value=report.lifecycle_years)
    ws.cell(row=meta_row + 4, column=1, value="CM-overlay PM uit").font = bold
    ws.cell(row=meta_row + 4, column=2, value=report.cm_overlay_pm_count)

    _autosize_columns(ws)

    legend = wb.create_sheet("Legenda")
    legend.cell(row=1, column=1, value="Kolom / term").font = bold
    legend.cell(row=1, column=2, value="Uitleg").font = bold
    for idx, (term, explanation) in enumerate(_LEGEND_ROWS, start=2):
        legend.cell(row=idx, column=1, value=term)
        legend.cell(row=idx, column=2, value=explanation)
    legend.column_dimensions["A"].width = 28
    legend.column_dimensions["B"].width = 72

    mark_start = len(_LEGEND_ROWS) + 3
    legend.cell(row=mark_start, column=1, value="Markeringen").font = bold
    legend.cell(row=mark_start, column=2, value="Kleur").font = bold
    legend.cell(row=mark_start, column=3, value="Betekenis").font = bold
    for offset, (label, color, meaning) in enumerate(_MARKING_LEGEND_ROWS, start=1):
        row_i = mark_start + offset
        legend.cell(row=row_i, column=1, value=label)
        swatch = legend.cell(row=row_i, column=2, value=" ")
        swatch.fill = PatternFill(fill_type="solid", fgColor=color)
        legend.cell(row=row_i, column=3, value=meaning)
    legend.column_dimensions["C"].width = 72

    causes = wb.create_sheet("Oorzaken")
    causes.cell(row=1, column=1, value="Code").font = bold
    causes.cell(row=1, column=2, value="Oorzaak").font = bold
    causes.cell(row=1, column=3, value="Actie").font = bold
    for idx, (code, label, action) in enumerate(_CAUSE_CATALOG_ROWS, start=2):
        causes.cell(row=idx, column=1, value=code)
        causes.cell(row=idx, column=2, value=label)
        causes.cell(row=idx, column=3, value=action)
    causes.column_dimensions["A"].width = 10
    causes.column_dimensions["B"].width = 42
    causes.column_dimensions["C"].width = 52

    if project is not None and fm_results is not None:
        _write_effect_categories_sheet(wb, project, fm_results)

    dest.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dest)
    return dest.resolve()


def _write_effect_categories_sheet(
    wb: Workbook,
    project: RCMProject,
    fm_results: Mapping[str, FMResult],
) -> None:
    ws = wb.create_sheet("Effectcategorieën")
    headers = (
        "FM",
        "Effectklasse",
        "Categorie",
        "RCM2-waarde",
        "Eenheid",
        "CM",
        "PM",
        "RF",
        "Oorzaakcode",
    )
    bold = Font(bold=True)
    for col_idx, header in enumerate(headers, start=1):
        ws.cell(row=1, column=col_idx, value=header).font = bold

    row_idx = 2
    for fm_id in sorted(fm_results):
        fmr = fm_results[fm_id]
        rows = aggregate(project, {fm_id: fmr}, fm_ids=frozenset({fm_id}))
        for er in rows:
            ws.cell(row=row_idx, column=1, value=fm_id)
            ws.cell(row=row_idx, column=2, value=er.label)
            ws.cell(row=row_idx, column=3, value=er.categorie)
            ws.cell(row=row_idx, column=4, value=er.waarde_totaal)
            ws.cell(row=row_idx, column=5, value=er.eenheid)
            ws.cell(row=row_idx, column=6, value=er.waarde_cm)
            ws.cell(row=row_idx, column=7, value=er.waarde_pm)
            ws.cell(row=row_idx, column=8, value=er.rf_display)
            ws.cell(row=row_idx, column=9, value="E2")
            row_idx += 1


def export_validation_excel(
    project: RCMProject,
    fm_results: Mapping[str, FMResult],
    path: Path | str,
    *,
    input_file: Path | str | None = None,
) -> Path:
    report = build_validation_report(project, fm_results)
    return export_validation_workbook(
        report,
        path,
        projectnaam=project.projectnaam,
        input_file=input_file,
        project=project,
        fm_results=fm_results,
    )


def validation_column_headers() -> Sequence[str]:
    return tuple(col.header for col in _VALIDATION_COLUMNS)
