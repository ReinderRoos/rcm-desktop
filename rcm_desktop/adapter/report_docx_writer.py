"""Schrijf ReportDocument naar Word (slice 57)."""

from __future__ import annotations

import io
from pathlib import Path

from rcm_desktop.adapter.report_document import (
    ReportDocument,
    ReportFigure,
    ReportSection,
    ReportTable,
)


def write_report_docx(document: ReportDocument, output_path: Path) -> Path:
    from docx import Document
    from docx.shared import Inches

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()

    def write_section(section: ReportSection, *, level: int) -> None:
        doc.add_heading(section.title, level=min(level, 3))
        for paragraph in section.paragraphs:
            doc.add_paragraph(paragraph.text)
        for table in section.tables:
            _write_table(doc, table)
        for figure in section.figures:
            _write_figure(doc, figure)
        for child in section.children:
            write_section(child, level=level + 1)

    for section in document.sections:
        write_section(section, level=1)

    doc.save(str(output_path))
    return output_path


def _write_table(doc, table: ReportTable) -> None:
    if not table.headers:
        return
    t = doc.add_table(rows=1, cols=len(table.headers))
    hdr = t.rows[0].cells
    for idx, header in enumerate(table.headers):
        hdr[idx].text = header
    for row in table.rows:
        cells = t.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = value


def _write_figure(doc, figure: ReportFigure) -> None:
    from docx.shared import Inches

    stream = io.BytesIO(figure.png_bytes)
    doc.add_picture(stream, width=Inches(6.0))
    if figure.caption:
        doc.add_paragraph(figure.caption)
